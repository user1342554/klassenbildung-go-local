from __future__ import annotations

import io

from openpyxl import load_workbook

from klassenbildung.core.models import ClassConfig, OptimizationSettings, ProfileSlackReport, SolverResult, Student
from klassenbildung.excel_io.excel_export import export_excel
from klassenbildung.optimization.solver import solve_assignments
from klassenbildung.presentation.candidate_summary import (
    CANDIDATE_SUMMARY_FIELDS,
    candidate_summary_records,
)
from klassenbildung.presentation.expert_result_view import expert_solver_diagnostic_text
from klassenbildung.presentation.standard_result_view import (
    STANDARD_MODE_FORBIDDEN_SOLVER_JARGON,
    all_candidate_summary_records_for_standard_view,
    standard_visible_text,
)
from klassenbildung.presentation.wording import candidate_tradeoff_text
from klassenbildung.services.candidate_selection import decision_candidate_cards, review_candidates


def test_candidate_cards_put_balanced_candidate_before_social_strongest() -> None:
    solver_result = SolverResult(
        "FEASIBLE",
        {},
        profile_slack_reports=[
            _report("C Musik +2", isolated=26, fl_actual=1, music_actual=3, fl_allowed=1, music_allowed=3),
            _report("E beide +1", isolated=28, fl_actual=2, music_actual=2, fl_allowed=2, music_allowed=2),
            _report("F mehr Profil-Slack", isolated=21, fl_actual=2, music_actual=3, fl_allowed=2, music_allowed=3),
        ],
    )

    cards = decision_candidate_cards(solver_result, 210)

    assert [summary.key for _, summary, _ in cards] == ["E", "F", "C"]
    assert cards[0][0] == "E - am ausgewogensten"


def test_e_wording_does_not_claim_socially_better_than_c_when_c_has_less_isolation() -> None:
    solver_result = SolverResult(
        "FEASIBLE",
        {},
        profile_slack_reports=[
            _report("C Musik +2", isolated=26, fl_actual=1, music_actual=3, fl_allowed=1, music_allowed=3),
            _report("E beide +1", isolated=28, fl_actual=2, music_actual=2, fl_allowed=2, music_allowed=2),
        ],
    )
    summaries = review_candidates(solver_result, 210)
    e_summary = next(summary for summary in summaries if summary.key == "E")

    text = candidate_tradeoff_text(e_summary, summaries)

    assert "Sozial liegt E knapp hinter C" in text
    assert "besser sozial" not in text


def test_candidate_summary_is_single_source_for_ui_json_export() -> None:
    solver_result = SolverResult(
        "FEASIBLE",
        {},
        profile_slack_reports=[
            _report(
                "A streng",
                isolated=58,
                fl_actual=1,
                music_actual=1,
                fl_allowed=1,
                music_allowed=1,
                fl_minority=12,
                music_minority=9,
                role=None,
            ),
            _report(
                "C Musik +2",
                isolated=26,
                fl_actual=1,
                music_actual=3,
                fl_allowed=1,
                music_allowed=3,
                fl_minority=18,
                music_minority=43,
                role="F/L-schonende Alternative",
            ),
            _report(
                "E beide +1",
                isolated=28,
                fl_actual=2,
                music_actual=2,
                fl_allowed=2,
                music_allowed=2,
                fl_minority=22,
                music_minority=24,
                role="Ausgewogenster Vorschlag",
            ),
            _report(
                "F mehr Profil-Slack",
                isolated=21,
                fl_actual=2,
                music_actual=3,
                fl_allowed=2,
                music_allowed=3,
                fl_minority=25,
                music_minority=45,
                role="Sozial stärkste Alternative",
            ),
        ],
    )

    json_records = _records_by_key(candidate_summary_records(solver_result, 210))
    ui_records = _records_by_key(all_candidate_summary_records_for_standard_view(solver_result, 210))
    excel_records = _records_by_key(_excel_candidate_records(solver_result))

    for key in ("A", "C", "E", "F"):
        for field in CANDIDATE_SUMMARY_FIELDS:
            assert ui_records[key][field] == json_records[key][field]
            assert excel_records[key][field] == json_records[key][field]


def test_debug_payload_uses_explicit_candidate_roles() -> None:
    payload = _debug_payload()

    assert payload["default_review_candidate"] == "E beide +1"
    assert payload["social_strongest_candidate"] == "F mehr Profil-Slack"
    assert payload["fl_preserving_candidate"] == "C Musik +2"
    assert "candidate_summaries" in payload
    assert "candidate_reviews" in payload


def test_debug_payload_does_not_emit_primary_recommendation() -> None:
    payload = _debug_payload()

    assert "primary_recommendation" not in payload


def test_debug_payload_does_not_emit_balanced_recommendation() -> None:
    payload = _debug_payload()

    assert "balanced_recommendation" not in payload


def test_debug_payload_does_not_emit_slack_candidates() -> None:
    payload = _debug_payload()

    assert "slack_candidates" not in payload


def test_current_debug_payload_does_not_emit_legacy_recommendation_keys_from_solver_run() -> None:
    import app

    students = [
        _student(1, "F", "B", friend1="2"),
        _student(2, "L", "S", friend1="1"),
        _student(3, "F", "G", friend1="4"),
        _student(4, "L", "B", friend1="3"),
    ]
    class_configs = [
        ClassConfig("5a", "5a", 2, 2, [], []),
        ClassConfig("5b", "5b", 2, 2, [], []),
    ]
    settings = OptimizationSettings(solver_time_limit_seconds=2)
    solver_result = solve_assignments(students, class_configs, settings)
    assert solver_result.score_report

    payload = app._solver_debug_payload(
        solver_result,
        students,
        class_configs,
        settings,
        solver_result.score_report,
    )
    keys = _nested_keys(payload)

    assert "primary_recommendation" not in keys
    assert "balanced_recommendation" not in keys
    assert "slack_candidates" not in keys
    assert "candidate_summaries" in payload
    assert "candidate_reviews" in payload
    assert "default_review_candidate" in payload
    assert "social_strongest_candidate" in payload
    assert "fl_preserving_candidate" in payload
    assert "overall_status" in payload


def test_current_candidate_payload_uses_explicit_role_names() -> None:
    payload = _debug_payload()

    for candidate in payload["candidate_summaries"]:
        keys = set(candidate)
        assert "role" in keys
        assert "recommendation_role" in keys
        assert "primary_recommendation" not in keys
        assert "balanced_recommendation" not in keys


def _debug_payload() -> dict[str, object]:
    import app

    solver_result = SolverResult(
        "FEASIBLE",
        {},
        profile_slack_reports=[
            _report("C Musik +2", isolated=26, fl_actual=1, music_actual=3, fl_allowed=1, music_allowed=3),
            _report("E beide +1", isolated=28, fl_actual=2, music_actual=2, fl_allowed=2, music_allowed=2),
            _report("F mehr Profil-Slack", isolated=21, fl_actual=2, music_actual=3, fl_allowed=2, music_allowed=3),
        ],
    )
    score = type("Score", (), {"total_score": 0, "hard_violations": [], "isolated_friend_request_count": 0})()

    return app._solver_debug_payload(solver_result, [], [], OptimizationSettings(), score)


def _nested_keys(value: object) -> set[str]:
    if isinstance(value, dict):
        keys = set(value)
        for item in value.values():
            keys.update(_nested_keys(item))
        return keys
    if isinstance(value, list):
        keys: set[str] = set()
        for item in value:
            keys.update(_nested_keys(item))
        return keys
    return set()


def test_standard_mode_hides_solver_jargon() -> None:
    solver_result = SolverResult(
        "FEASIBLE",
        {},
        profile_slack_reports=[
            _report("A streng", isolated=58, fl_actual=1, music_actual=1, fl_allowed=1, music_allowed=1),
            _report("E beide +1", isolated=28, fl_actual=2, music_actual=2, fl_allowed=2, music_allowed=2),
            _report("F mehr Profil-Slack", isolated=21, fl_actual=2, music_actual=3, fl_allowed=2, music_allowed=3),
        ],
    )

    text = standard_visible_text(solver_result, 210)

    for term in STANDARD_MODE_FORBIDDEN_SOLVER_JARGON:
        assert term not in text


def test_expert_mode_shows_solver_diagnostics() -> None:
    text = expert_solver_diagnostic_text()

    for term in ("FEASIBLE", "UNKNOWN", "OPTIMAL", "Gap", "Objective", "Best Bound"):
        assert term in text


def _report(
    variant: str,
    *,
    isolated: int,
    fl_actual: int,
    music_actual: int,
    fl_allowed: int,
    music_allowed: int,
    fl_minority: int = 0,
    music_minority: int = 0,
    role: str | None = None,
) -> ProfileSlackReport:
    return ProfileSlackReport(
        variant=variant,
        language_mixed_limit=fl_allowed,
        music_mixed_limit=music_allowed,
        status="FEASIBLE",
        mixed_language_class_count=fl_actual,
        mixed_music_class_count=music_actual,
        language_minority_student_count=fl_minority,
        music_minority_student_count=music_minority,
        isolated_friend_request_count=isolated,
        friend1_fulfilled=160,
        friend1_total=210,
        friend2_fulfilled=120,
        friend2_total=180,
        mutual_friend_fulfilled=80,
        mutual_friend_total=105,
        review_candidate=variant != "A streng",
        social_limit_met=True,
        recommendation_role=role,
        assignments={f"s{index}": f"{variant}-5a" for index in range(210)},
    )


def _student(index: int, language: str, music: str = "Reg", *, friend1: str | None = None) -> Student:
    return Student(
        internal_id=f"s{index}",
        row_number=index,
        original_class=None,
        nr=str(index),
        school="Grundschule",
        last_name=f"N{index}",
        first_name=f"V{index}",
        eligibility="GYM",
        gender="w" if index % 2 else "m",
        birthdate=None,
        nationality="DE",
        religion="ev",
        second_language=language,
        music_profile=music,
        primary_class="4a",
        friend1=friend1,
        friend2=None,
        comment=None,
        note_text=None,
        is_support=False,
    )


def _records_by_key(records: list[dict[str, object]]) -> dict[str, dict[str, object]]:
    return {str(record["key"]): record for record in records}


def _excel_candidate_records(solver_result: SolverResult) -> list[dict[str, object]]:
    exported = export_excel(
        None,
        [],
        {},
        [],
        None,
        [],
        profile_slack_reports=solver_result.profile_slack_reports,
    )
    workbook = load_workbook(io.BytesIO(exported))
    rows = list(workbook["Kandidaten"].iter_rows(values_only=True))
    header = list(rows[0])
    return [dict(zip(header, row, strict=True)) for row in rows[1:]]
