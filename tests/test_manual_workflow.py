from __future__ import annotations

import io

from openpyxl import load_workbook

import klassenbildung.optimization.solver as solver_module
from klassenbildung.core.models import ClassConfig, ManualRule, OptimizationSettings, Student
from klassenbildung.excel_io.excel_export import export_excel
from klassenbildung.optimization.solver import solve_assignments
from klassenbildung.presentation.candidate_review import build_candidate_review_model
from klassenbildung.presentation.result_view_model import CandidateRole, CandidateSource, CandidateSummary
from klassenbildung.services.manual_rules import (
    NoteReviewStatus,
    active_manual_rules,
    add_manual_rule_entry,
    manual_rule_entry_records,
    note_review_status_by_student,
)
from klassenbildung.services.note_rule_conversion import convert_note_to_manual_rule


def test_full_manual_workflow_note_rule_solve_review_export() -> None:
    students = [
        _student(1, "F", "B", friend1="2", note_text="nicht mit V2 zusammen"),
        _student(2, "F", "B", friend1="1"),
        _student(3, "L", "S", friend1="4"),
        _student(4, "L", "S", friend1="3"),
    ]
    class_configs = [
        ClassConfig("5a", "5a", 0, 4, [], []),
        ClassConfig("5b", "5b", 0, 4, [], []),
    ]
    settings = OptimizationSettings(solver_time_limit_seconds=2)

    conversion = convert_note_to_manual_rule(
        students,
        "s1",
        "SEPARATE",
        selected_student_id="s2",
        confirmed=True,
    )
    entries, changed = add_manual_rule_entry([], conversion.rule, source="note", note_student_id="s1")

    assert changed
    assert manual_rule_entry_records(entries, students)[0]["Status"] == "aktiv"

    base_cache_key = solver_module._profile_incumbent_cache_key(students, class_configs, settings, [])
    rule_cache_key = solver_module._profile_incumbent_cache_key(
        students,
        class_configs,
        settings,
        active_manual_rules(entries),
    )
    assert rule_cache_key != base_cache_key

    solver_result = solve_assignments(students, class_configs, settings, manual_rules=active_manual_rules(entries))

    assert solver_result.assignments["s1"] != solver_result.assignments["s2"]
    assert solver_result.score_report is not None
    assert not solver_result.score_report.hard_violations

    summary = _summary_from_solver_result(solver_result.assignments, solver_result.score_report, len(students))
    note_statuses = note_review_status_by_student(entries, set())
    review = build_candidate_review_model(
        summary,
        students,
        class_configs,
        settings,
        note_review_status_by_student=note_statuses,
    )

    assert review.students_with_manual_notes[0].review_status == NoteReviewStatus.CONVERTED_TO_RULE

    exported = export_excel(
        None,
        students,
        solver_result.assignments,
        class_configs,
        solver_result.score_report,
        [],
        manual_rule_entries=entries,
        note_review_status_by_student=note_statuses,
        settings=settings,
    )
    workbook = load_workbook(io.BytesIO(exported))

    assert workbook["Manuelle Regeln"]["A2"].value == "aktiv"
    assert workbook["Notizen"]["C2"].value == "in Regel umgewandelt"
    assert "Manuelle Änderungen" not in workbook.sheetnames
    overview = {
        row[0]: row[1]
        for row in workbook["Übersicht"].iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    assert overview["Anzahl aktiver Regeln"] == 1
    assert overview["Anzahl ungeprüfter Notizen"] == 0


def _summary_from_solver_result(assignments: dict[str, str], score, student_count: int) -> CandidateSummary:
    return CandidateSummary(
        key="E",
        name="E beide +1",
        role=CandidateRole.REVIEW,
        recommendation_role="Ausgewogenster Vorschlag",
        assignments=assignments,
        fl_mixed_actual=score.mixed_language_class_count,
        fl_mixed_allowed=2,
        music_mixed_actual=score.mixed_music_class_count,
        music_mixed_allowed=2,
        fl_minority=score.language_minority_student_count,
        music_minority=score.music_minority_student_count,
        without_wishfriend=score.isolated_friend_request_count,
        without_wishfriend_rate=score.isolated_friend_request_count / max(student_count, 1),
        friend1_satisfied=score.friend1_fulfilled,
        friend1_total=score.friend1_total,
        mutual_satisfied=score.mutual_friend_fulfilled,
        mutual_total=score.mutual_friend_total,
        friend2_satisfied=score.friend2_fulfilled,
        friend2_total=score.friend2_total,
        social_threshold_met=True,
        candidate_for_review=True,
        automatically_approvable=False,
        gap=None,
        gap_reliable=False,
        displayed_source=CandidateSource.SLACK_SEARCH,
        refinement_attempted=False,
        refinement_status=None,
        refinement_result=None,
    )


def _student(
    index: int,
    language: str,
    music: str,
    *,
    friend1: str | None = None,
    note_text: str | None = None,
) -> Student:
    return Student(
        internal_id=f"s{index}",
        row_number=index,
        original_class=None,
        nr=str(index),
        school="Grundschule",
        last_name=f"N{index}",
        first_name=f"V{index}",
        eligibility="GYM",
        gender="w" if index % 2 else "m",
        birthdate=None,
        nationality="DE",
        religion="ev",
        second_language=language,
        music_profile=music,
        primary_class="4a",
        friend1=friend1,
        friend2=None,
        comment=note_text,
        note_text=note_text,
        is_support=False,
    )
